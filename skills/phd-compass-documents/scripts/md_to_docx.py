#!/usr/bin/env python3
"""Build a clean Word document (.docx) from a small Markdown dialect. Needs python-docx.

Usage: python3 md_to_docx.py <input.md> <output.docx>

Dialect (one construct per line):
  <!-- docx: paper=A4|Letter margin_cm=2.2 font=Calibri size=10.5 -->   optional first line
  # Name                         document title (centred, large)
  plain lines right after the title   header lines (centred, small), e.g. contact details
  ## Section                     section heading with a rule underneath
  ### Title | Place | Dates      entry: bold title, right-aligned dates, italic place on the next line
                                  (two fields = title and dates)
  - item / "  - sub-item"         bullets
  blank line                     paragraph break
  other text                     paragraph
  Inline: **bold**, *italic*, [text](https://link)
"""

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    print("python-docx is not installed. Install it with: pip install python-docx", file=sys.stderr)
    sys.exit(2)

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")
MUTED = RGBColor(0x55, 0x55, 0x55)


def options(first_line):
    opts = {"paper": "A4", "margin_cm": "2.2", "font": "Calibri", "size": "10.5"}
    m = re.match(r"<!--\s*docx:(.*?)-->", first_line.strip())
    if m:
        for k, v in re.findall(r"(\w+)=([^\s]+)", m.group(1)):
            opts[k] = v
    return opts, bool(m)


def add_hyperlink(paragraph, text, url, size=None):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    props.append(color)
    under = OxmlElement("w:u")
    under.set(qn("w:val"), "single")
    props.append(under)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        props.append(sz)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text, size=None, color=None, italic=False, bold=False):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
            run.italic = italic
        elif piece.startswith("*"):
            run = paragraph.add_run(piece[1:-1])
            run.italic = True
            run.bold = bold
        elif piece.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece)
            add_hyperlink(paragraph, m.group(1), m.group(2), size)
            continue
        else:
            run = paragraph.add_run(piece)
            run.italic = italic
            run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def rule_below(paragraph):
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def build(src, out):
    lines = Path(src).read_text(encoding="utf-8").splitlines()
    opts, had_opts = options(lines[0] if lines else "")
    if had_opts:
        lines = lines[1:]
    size = float(opts["size"])

    doc = Document()
    sec = doc.sections[0]
    if opts["paper"].lower() == "letter":
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    margin = Cm(float(opts["margin_cm"]))
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = margin
    text_width = sec.page_width - sec.left_margin - sec.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = opts["font"]
    normal.font.size = Pt(size)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), opts["font"])
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.1

    in_header = False
    buffer = []

    def flush():
        if buffer:
            p = doc.add_paragraph()
            add_inline(p, " ".join(s.strip() for s in buffer))
            buffer.clear()

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("# "):
            flush()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.size = Pt(size + 8)
            in_header = True
            continue
        if in_header and line and not line.startswith(("#", "-")):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, line.strip(), size=size - 1, color=MUTED)
            continue
        if in_header and not line:
            in_header = False
            continue
        in_header = False
        if line.startswith("## "):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[3:].strip())
            run.bold = True
            run.font.size = Pt(size + 2)
            rule_below(p)
        elif line.startswith("### "):
            flush()
            fields = [f.strip() for f in line[4:].split("|")]
            title, place, dates = (fields + ["", ""])[:3] if len(fields) != 2 else (fields[0], "", fields[1])
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
            add_inline(p, title, bold=True)
            if dates:
                run = p.add_run("\t" + dates)
                run.font.size = Pt(size - 0.5)
            if place:
                q = doc.add_paragraph()
                q.paragraph_format.space_after = Pt(1)
                q.paragraph_format.keep_with_next = True
                add_inline(q, place, size=size - 0.5, italic=True)
        elif re.match(r"^\s*- ", line):
            flush()
            level = 2 if re.match(r"^\s{2,}- ", line) else 1
            p = doc.add_paragraph(style="List Bullet" if level == 1 else "List Bullet 2")
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, re.sub(r"^\s*- ", "", line))
        elif not line.strip():
            flush()
        else:
            buffer.append(line)
    flush()
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    build(sys.argv[1], sys.argv[2])
